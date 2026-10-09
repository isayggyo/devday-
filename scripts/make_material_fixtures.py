"""Author a real two-slide PPTX fixture; never inject product AI output."""
from pathlib import Path
from pptx import Presentation

root = Path(__file__).resolve().parents[1]
deck = Presentation()
for title, text in [
    ("Retrieval practice", "Recall the lecture from memory before reviewing notes.\nThe act of retrieval strengthens learning."),
    ("Spaced practice", "Review the lecture across separate study sessions.\nCombine retrieval practice with spacing for long-term retention."),
]:
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = text
deck.save(root / "e2e/fixtures/lecture.pptx")
print("Authored PPTX fixture: 2 slides")
